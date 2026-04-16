"""
PPT Generator Service
Generates professional PowerPoint presentations for properties using python-pptx
"""
import io
import random
import time
import logging
import math
from typing import Optional, List
from datetime import datetime

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import httpx

from app.services.ppt.templates import TEMPLATES, PPTTemplate

logger = logging.getLogger(__name__)


class PropertyPPTGenerator:
    """
    Generates professional PowerPoint presentations for properties.
    A random visual template is chosen on each instantiation so every
    export looks unique while carrying the same data.
    """

    # Slide dimensions (16:9 widescreen)
    SLIDE_WIDTH = Inches(13.333)
    SLIDE_HEIGHT = Inches(7.5)

    def __init__(self, template_name: str = None):
        """
        Initialize the PPT generator.

        Args:
            template_name: Name of the template to use (from PPTTemplate.name).
                           If None or not found, a random template is chosen.
        """
        self.prs: Optional[Presentation] = None
        if template_name:
            matched = next((t for t in TEMPLATES if t.name == template_name), None)
            self.template: PPTTemplate = matched or random.choice(TEMPLATES)
        else:
            self.template: PPTTemplate = random.choice(TEMPLATES)

    def generate(self, property_data: dict, images: List[bytes] = None) -> io.BytesIO:
        """
        Generate a complete PPT presentation for a property.

        Args:
            property_data: Dictionary containing property details
            images: List of image bytes (already optimized)

        Returns:
            BytesIO object containing the PPTX file
        """
        start_time = time.time()

        # Create new presentation (16:9 widescreen)
        self.prs = Presentation()
        self.prs.slide_width = self.SLIDE_WIDTH
        self.prs.slide_height = self.SLIDE_HEIGHT

        # Generate slides
        self._create_title_slide(property_data, images[0] if images else None)
        self._create_overview_slide(property_data)
        self._create_details_slide(property_data)
        self._create_features_slide(property_data)

        if images and len(images) > 1:
            self._create_gallery_slide(images[1:5])  # Max 4 gallery images

        if property_data.get("latitude") and property_data.get("longitude"):
            self._create_location_slide(property_data)

        self._create_contact_slide(property_data)

        # Save to BytesIO
        output = io.BytesIO()
        self.prs.save(output)
        output.seek(0)

        duration_ms = int((time.time() - start_time) * 1000)
        logger.info(
            f"PPT generated in {duration_ms}ms for property "
            f"{property_data.get('our_id')} [template={self.template.name}]"
        )

        return output

    # ── Helpers ──────────────────────────────────────────────────

    def _apply_background(self, slide, color: RGBColor = None):
        """Apply a solid background color to a slide using the slide background API."""
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = color or self.template.slide_bg

    def _add_divider(self, slide, top_inches: float = 1):
        """Add a horizontal divider line."""
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(top_inches), Inches(12.3), Inches(0.02)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = self.template.accent
        line.line.fill.background()
        return line

    def _add_shape_with_text(
        self,
        slide,
        left: float,
        top: float,
        width: float,
        height: float,
        text: str,
        font_size: int = 14,
        font_color: RGBColor = None,
        bold: bool = False,
        alignment: PP_ALIGN = PP_ALIGN.LEFT,
        fill_color: RGBColor = None,
    ):
        """Helper to add a text box with styling"""
        shape = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = None

        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.name = self.template.font_name
        p.font.bold = bold
        p.font.color.rgb = font_color or self.template.body_text
        p.alignment = alignment

        if fill_color:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_color

        return shape

    def _add_rounded_rectangle(
        self,
        slide,
        left: float,
        top: float,
        width: float,
        height: float,
        fill_color: RGBColor,
        text: str = None,
        font_size: int = 12,
        font_color: RGBColor = None,
    ):
        """Add a rounded rectangle shape"""
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left), Inches(top), Inches(width), Inches(height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
        shape.line.fill.background()  # No border

        if text:
            tf = shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = text
            p.font.size = Pt(font_size)
            p.font.name = self.template.font_name
            p.font.color.rgb = font_color or self.template.white
            p.alignment = PP_ALIGN.CENTER
            tf.anchor = MSO_ANCHOR.MIDDLE

        return shape

    def _format_heading(self, text: str) -> str:
        """Apply template-specific heading formatting (e.g. uppercase)."""
        return text.upper() if self.template.heading_uppercase else text

    def _add_card_border(self, slide, left: float, top: float,
                         width: float, height: float):
        """
        Draw the template-specific decorative border for metric / section cards.
        - Dark Luxury: full border drawn via shape.line
        - Clean Minimal: 3 pt blue left-side bar
        - Modern Gradient: 3 pt teal top bar
        Returns the border shape (or None).
        """
        t = self.template
        if t.card_left_border and t.card_border_color:
            # Thin rectangle on the left edge
            bar_w = t.card_border_width_pt / 72  # convert pt to inches
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(left), Inches(top),
                Inches(bar_w), Inches(height),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = t.card_border_color
            bar.line.fill.background()
            return bar

        if t.card_top_border and t.card_border_color:
            bar_h = t.card_border_width_pt / 72
            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(left), Inches(top),
                Inches(width), Inches(bar_h),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = t.card_border_color
            bar.line.fill.background()
            return bar

        return None

    def _style_card_shape(self, shape):
        """Apply full-border styling to a card shape (Dark Luxury pattern)."""
        t = self.template
        if (
            t.card_border_color
            and t.card_border_width_pt > 0
            and not t.card_left_border
            and not t.card_top_border
        ):
            shape.line.color.rgb = t.card_border_color
            shape.line.width = Pt(t.card_border_width_pt)

    # ── Slide builders ──────────────────────────────────────────

    def _create_title_slide(self, property_data: dict, hero_image: bytes = None):
        """Create the title/cover slide"""
        t = self.template
        slide_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(slide_layout)

        # Slide background
        self._apply_background(slide)

        # Text panel background rectangle
        panel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(t.title_panel_left), Inches(t.title_panel_top),
            Inches(t.title_panel_width), Inches(t.title_panel_height),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = t.title_panel_color
        panel.line.fill.background()

        # Optional accent line at the top of the text panel (Modern Gradient)
        if t.title_accent_line:
            accent_h = t.title_accent_line_height_pt / 72
            accent_line = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(t.title_panel_left), Inches(t.title_panel_top),
                Inches(t.title_panel_width), Inches(accent_h),
            )
            accent_line.fill.solid()
            accent_line.fill.fore_color.rgb = t.accent
            accent_line.line.fill.background()

        # Hero image (if available)
        if hero_image:
            try:
                img_stream = io.BytesIO(hero_image)
                slide.shapes.add_picture(
                    img_stream,
                    Inches(t.title_hero_left), Inches(t.title_hero_top),
                    width=Inches(t.title_hero_width),
                    height=Inches(t.title_hero_height),
                )
            except Exception as e:
                logger.warning(f"Failed to add hero image: {e}")

        # Property title — use template font size and a tall enough box for long titles
        title = property_data.get("title", "Property Presentation")
        self._add_shape_with_text(
            slide, t.title_text_left, t.title_text_top,
            t.title_text_width, t.title_text_height,
            title,
            font_size=t.title_font_size, font_color=t.body_text, bold=True
        )

        # Location — placed at the template's safe absolute Y so it never overlaps the title
        location = property_data.get("area_name", "")
        if location:
            self._add_shape_with_text(
                slide, t.title_text_left, t.title_loc_top,
                t.title_text_width, 0.6,
                f"📍 {location}",
                font_size=14, font_color=t.accent, bold=False
            )

        # Price badge — placed at its own safe absolute Y
        price = property_data.get("current_price")
        currency = property_data.get("currency", "PKR")
        if price:
            formatted_price = self._format_price(price, currency)
            self._add_rounded_rectangle(
                slide, t.title_text_left, t.title_price_top,
                2.5, 0.55,
                t.accent,
                formatted_price,
                font_size=16, font_color=t.white
            )

        # PropertyWalay branding — fixed at very bottom of slide, small and unobtrusive
        self._add_shape_with_text(
            slide, 0.5, 7.1, 5, 0.35,
            "PropertyWalay - AI-Powered Real Estate",
            font_size=9, font_color=t.muted_text
        )

        # Date — bottom-right corner
        self._add_shape_with_text(
            slide, 9.5, 7.1, 3.5, 0.35,
            datetime.now().strftime("%B %d, %Y"),
            font_size=9, font_color=t.muted_text,
            alignment=PP_ALIGN.RIGHT
        )

    def _create_overview_slide(self, property_data: dict):
        """Create property overview slide with key stats"""
        t = self.template
        slide_layout = self.prs.slide_layouts[6]  # Blank
        slide = self.prs.slides.add_slide(slide_layout)

        # Background
        self._apply_background(slide)

        # Title
        self._add_shape_with_text(
            slide, 0.5, 0.3, 12, 0.8,
            self._format_heading("Property Overview"),
            font_size=28, bold=True, font_color=t.accent
        )

        # Horizontal line
        self._add_divider(slide, top_inches=1)

        # Key stats boxes
        stats = self._get_property_stats(property_data)
        box_width = 2.8
        box_height = 1.5
        start_x = 0.5
        gap = 0.3

        for i, (label, value, icon) in enumerate(stats[:4]):
            x = start_x + i * (box_width + gap)

            # Stat box
            box = self._add_rounded_rectangle(
                slide, x, 1.5, box_width, box_height,
                t.card_bg
            )
            self._style_card_shape(box)

            # Decorative per-side border
            self._add_card_border(slide, x, 1.5, box_width, box_height)

            # Icon and value
            self._add_shape_with_text(
                slide, x + 0.2, 1.7, box_width - 0.4, 0.5,
                f"{icon} {value}",
                font_size=20, bold=True, font_color=t.card_number_color
            )

            # Label
            self._add_shape_with_text(
                slide, x + 0.2, 2.3, box_width - 0.4, 0.4,
                label,
                font_size=12, font_color=t.card_label_color
            )

        # Property description
        self._add_shape_with_text(
            slide, 0.5, 3.5, 12.3, 0.5,
            self._format_heading("Description"),
            font_size=18, bold=True, font_color=t.accent
        )

        description = self._generate_description(property_data)
        self._add_shape_with_text(
            slide, 0.5, 4.1, 12.3, 2.5,
            description,
            font_size=14, font_color=t.muted_text
        )

    def _create_details_slide(self, property_data: dict):
        """Create detailed property information slide"""
        t = self.template
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Background
        self._apply_background(slide)

        # Title
        self._add_shape_with_text(
            slide, 0.5, 0.3, 12, 0.8,
            self._format_heading("Property Details"),
            font_size=28, bold=True, font_color=t.accent
        )

        # Horizontal line
        self._add_divider(slide, top_inches=1)

        # Two columns of details
        details_left = [
            ("Property Type", property_data.get("prop_type", "N/A")),
            ("Subtype", property_data.get("prop_subtype", "N/A")),
            ("Bedrooms", str(property_data.get("beds", "N/A"))),
            ("Bathrooms", str(property_data.get("baths", "N/A"))),
        ]

        area_size = property_data.get("area_size")
        area_unit = property_data.get("area_unit", "")
        area_display = f"{area_size} {area_unit}" if area_size else "N/A"

        price = property_data.get("current_price")
        currency = property_data.get("currency", "PKR")
        price_display = self._format_price(price, currency) if price else "N/A"

        details_right = [
            ("Area Size", area_display),
            ("Price", price_display),
            ("Source", property_data.get("source", "N/A").capitalize()),
            ("Property ID", property_data.get("source_human_id") or str(property_data.get("our_id", ""))[:8]),
        ]

        # Left column
        y = 1.5
        for label, value in details_left:
            self._add_shape_with_text(
                slide, 0.5, y, 2.5, 0.4,
                label,
                font_size=12, font_color=t.muted_text
            )
            self._add_shape_with_text(
                slide, 3.2, y, 3, 0.4,
                str(value),
                font_size=14, bold=True, font_color=t.body_text
            )
            y += 0.7

        # Right column
        y = 1.5
        for label, value in details_right:
            self._add_shape_with_text(
                slide, 7, y, 2.5, 0.4,
                label,
                font_size=12, font_color=t.muted_text
            )
            self._add_shape_with_text(
                slide, 9.7, y, 3, 0.4,
                str(value),
                font_size=14, bold=True, font_color=t.body_text
            )
            y += 0.7

        # Price per unit (if applicable)
        if price and area_size:
            price_per_unit = price / area_size
            self._add_shape_with_text(
                slide, 0.5, 5, 6, 0.5,
                f"Price per {area_unit}: {self._format_price(price_per_unit, currency)}",
                font_size=14, font_color=t.accent
            )

    def _create_features_slide(self, property_data: dict):
        """
        Create features/amenities slide if scraped data is available.
        """
        t = self.template
        scraped_features = property_data.get("scraped_features") or {}
        scraped_amenities = property_data.get("scraped_amenities") or {}

        # Build grouped sections
        sections: list[tuple[str, list[str]]] = []

        for label, items in scraped_features.items():
            if items:
                sections.append((f"{label} Features", items))

        for label, items in scraped_amenities.items():
            if items:
                sections.append((label, items))

        if not sections:
            return  # Nothing to render

        # Chunk sections into pages of 2 sections per slide (1x2 layout)
        per_slide = 2
        chunks = [sections[i:i + per_slide] for i in range(0, len(sections), per_slide)]

        for chunk in chunks:
            slide_layout = self.prs.slide_layouts[6]
            slide = self.prs.slides.add_slide(slide_layout)

            # Background
            self._apply_background(slide)

            # Title
            self._add_shape_with_text(
                slide, 0.5, 0.3, 12, 0.8,
                self._format_heading("Features & Amenities"),
                font_size=28, bold=True, font_color=t.accent
            )

            # Horizontal line
            self._add_divider(slide, top_inches=1)

            # Layout: 2 cards stacked vertically, full width
            x = 0.5
            col_width = 12.3
            y_start = 1.3
            card_height = 2.8
            gap_y = 0.4

            for i, (title, items) in enumerate(chunk):
                y = y_start + i * (card_height + gap_y)

                # Card background
                card = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    Inches(x),
                    Inches(y),
                    Inches(col_width),
                    Inches(card_height),
                )
                card.fill.solid()
                card.fill.fore_color.rgb = t.card_bg
                card.line.fill.background()
                self._style_card_shape(card)

                # Decorative border
                self._add_card_border(slide, x, y, col_width, card_height)

                # Section title
                self._add_shape_with_text(
                    slide, x + 0.2, y + 0.2, col_width - 0.4, 0.4,
                    title,
                    font_size=16, bold=True, font_color=t.body_text
                )

                # Bullet list (cap to 10 items to avoid overflow)
                bullets_text = "\n".join([f"• {item}" for item in items[:10]])
                self._add_shape_with_text(
                    slide, x + 0.2, y + 0.7, col_width - 0.4, card_height - 0.9,
                    bullets_text,
                    font_size=12, font_color=t.muted_text
                )

    def _create_gallery_slide(self, images: List[bytes]):
        """Create image gallery slide"""
        t = self.template
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Background
        self._apply_background(slide)

        # Title
        self._add_shape_with_text(
            slide, 0.5, 0.3, 12, 0.8,
            self._format_heading("Property Gallery"),
            font_size=28, bold=True, font_color=t.accent
        )

        # Grid of images (2x2)
        positions = [
            (0.5, 1.2, 6, 3),
            (6.8, 1.2, 6, 3),
            (0.5, 4.4, 6, 3),
            (6.8, 4.4, 6, 3),
        ]

        for i, img_bytes in enumerate(images[:4]):
            if i >= len(positions):
                break
            try:
                left, top, width, height = positions[i]
                img_stream = io.BytesIO(img_bytes)
                slide.shapes.add_picture(
                    img_stream,
                    Inches(left), Inches(top),
                    width=Inches(width), height=Inches(height)
                )
            except Exception as e:
                logger.warning(f"Failed to add gallery image {i}: {e}")

    def _create_location_slide(self, property_data: dict):
        """Create location slide with map placeholder"""
        t = self.template
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Background
        self._apply_background(slide)

        # Title
        self._add_shape_with_text(
            slide, 0.5, 0.3, 12, 0.8,
            self._format_heading("Location"),
            font_size=28, bold=True, font_color=t.accent
        )

        # Location details
        location = property_data.get("area_name", "Location not specified")
        self._add_shape_with_text(
            slide, 0.5, 1.2, 12, 0.5,
            f"📍 {location}",
            font_size=18, font_color=t.accent
        )

        # Map or placeholder
        lat = property_data.get("latitude")
        lng = property_data.get("longitude")
        map_left, map_top = Inches(0.5), Inches(2)
        map_width, map_height = Inches(12.3), Inches(5)

        map_image = None
        if lat is not None and lng is not None:
            try:
                # Convert to float if needed (handles Decimal, string, etc.)
                lat_float = float(lat)
                lng_float = float(lng)
                logger.info(f"Fetching map for coordinates: {lat_float}, {lng_float}")
                map_image = self._fetch_static_map(lat_float, lng_float)
                if map_image:
                    logger.info(f"Map image fetched successfully: {len(map_image)} bytes")
            except (ValueError, TypeError) as e:
                logger.warning(f"Invalid coordinate format: lat={lat}, lng={lng}, error={e}")
            except Exception as e:
                logger.warning(f"Map fetch failed: {type(e).__name__}: {e}")

        if map_image:
            try:
                slide.shapes.add_picture(
                    io.BytesIO(map_image),
                    map_left,
                    map_top,
                    width=map_width,
                    height=map_height,
                )
                logger.info("Map image added to slide successfully")
            except Exception as e:
                logger.error(f"Failed to add map image to slide: {e}")
                # Fall through to placeholder
                map_image = None

        if not map_image:
            # Show placeholder with coordinates if available
            map_placeholder = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                map_left,
                map_top,
                map_width,
                map_height,
            )
            map_placeholder.fill.solid()
            map_placeholder.fill.fore_color.rgb = t.card_bg
            map_placeholder.line.color.rgb = RGBColor(55, 65, 81)  # subtle border

            if lat is not None and lng is not None:
                try:
                    lat_float = float(lat)
                    lng_float = float(lng)
                    coord_text = f"Coordinates:\n{lat_float:.6f}, {lng_float:.6f}\n\nMap unavailable"
                except (ValueError, TypeError):
                    coord_text = "Map not available"
            else:
                coord_text = "Map not available\n(Coordinates not provided)"

            self._add_shape_with_text(
                slide, 5, 3.8, 4, 1.4,
                coord_text,
                font_size=14, font_color=t.muted_text,
                alignment=PP_ALIGN.CENTER
            )

    def _create_contact_slide(self, property_data: dict):
        """Create contact/CTA slide"""
        t = self.template
        slide_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(slide_layout)

        # Background
        self._apply_background(slide)

        # Title
        self._add_shape_with_text(
            slide, 0.5, 1.5, 12.3, 1,
            "Interested in this property?",
            font_size=36, bold=True, font_color=t.body_text,
            alignment=PP_ALIGN.CENTER
        )

        # Contact info
        poc_name = property_data.get("poc_name") or "Contact Agent"
        poc_number = property_data.get("poc_number") or "Phone not provided"

        self._add_shape_with_text(
            slide, 0.5, 3, 12.3, 0.6,
            f"Agent: {poc_name}",
            font_size=20, font_color=t.body_text,
            alignment=PP_ALIGN.CENTER
        )

        if poc_number:
            self._add_shape_with_text(
                slide, 0.5, 3.7, 12.3, 0.6,
                f"📞 {poc_number}",
                font_size=18, font_color=t.accent,
                alignment=PP_ALIGN.CENTER
            )

        # CTA Button
        self._add_rounded_rectangle(
            slide, 4.5, 5, 4.3, 0.8,
            t.accent,
            "Schedule a Visit",
            font_size=18, font_color=t.white
        )

        # Footer
        self._add_shape_with_text(
            slide, 0.5, 6.8, 12.3, 0.4,
            "Generated by PropertyWalay - Your AI-Powered Real Estate Assistant",
            font_size=10, font_color=t.muted_text,
            alignment=PP_ALIGN.CENTER
        )

    # ── Data helpers (unchanged) ─────────────────────────────────

    def _get_property_stats(self, property_data: dict) -> List[tuple]:
        """Get key property stats for overview with sensible fallbacks."""
        stats = []

        beds = property_data.get("beds")
        baths = property_data.get("baths")
        area_size = property_data.get("area_size")
        area_unit = property_data.get("area_unit", "")
        prop_type = property_data.get("prop_type")
        price = property_data.get("current_price")
        currency = property_data.get("currency", "PKR")
        source = property_data.get("source")

        if beds is not None:
            stats.append(("Bedrooms", str(beds), "🛏️"))

        if baths is not None:
            stats.append(("Bathrooms", str(baths), "🚿"))

        if area_size:
            stats.append(("Area", f"{area_size} {area_unit}".strip(), "📐"))

        if prop_type:
            stats.append(("Type", prop_type.capitalize(), "🏠"))

        # Fallbacks when common fields are missing
        if price:
            stats.append(("Price", self._format_price(price, currency), "💰"))

        if source:
            stats.append(("Source", source.capitalize(), "🔗"))

        # Ensure exactly 4 entries with friendly placeholders
        while len(stats) < 4:
            stats.append(("Info", "Not available", "ℹ️"))

        return stats[:4]

    def _generate_description(self, property_data: dict) -> str:
        """Generate a property description resilient to missing fields."""
        parts = []

        prop_type = property_data.get("prop_type", "property") or "property"
        prop_subtype = property_data.get("prop_subtype") or ""

        # Insert space in camelCase subtype like residentialPlot -> residential Plot
        if prop_subtype and prop_subtype != prop_subtype.lower():
            prop_subtype = "".join(
                (" " + c if c.isupper() else c) for c in prop_subtype
            ).strip()

        beds = property_data.get("beds")
        baths = property_data.get("baths")
        area_size = property_data.get("area_size")
        area_unit = property_data.get("area_unit", "")
        location = property_data.get("area_name", "")
        source = property_data.get("source", "")

        type_str = f"{prop_subtype} {prop_type}".strip() if prop_subtype else prop_type
        parts.append(f"This {type_str.lower()}")

        if beds and baths:
            parts.append(f"features {beds} bedroom(s) and {baths} bathroom(s)")
        elif beds:
            parts.append(f"features {beds} bedroom(s)")
        elif baths:
            parts.append(f"features {baths} bathroom(s)")

        if area_size:
            parts.append(f"with a total area of {area_size} {area_unit}".strip())

        if location:
            parts.append(f"located in {location}")

        description = " ".join(parts).strip()
        if description and not description.endswith("."):
            description += "."

        if source:
            description += f" Listed on {source.capitalize()}."

        return description or "Property details are not available."

    def _format_price(self, price: float, currency: str = "PKR") -> str:
        """Format price with currency"""
        if price >= 10000000:  # 1 Crore
            return f"{currency} {price / 10000000:.2f} Cr"
        elif price >= 100000:  # 1 Lac
            return f"{currency} {price / 100000:.2f} Lac"
        else:
            return f"{currency} {price:,.0f}"

    def _fetch_static_map(self, lat: float, lng: float) -> Optional[bytes]:
        """
        Fetch a static map image using multiple fallback services.
        Tries OpenStreetMap, then alternative services.
        Returns image bytes or None on failure.
        """
        # Ensure coordinates are floats and valid
        try:
            lat = float(lat)
            lng = float(lng)

            # Validate coordinates (rough bounds for Pakistan, but allow wider range)
            if not (-90.0 <= lat <= 90.0) or not (-180.0 <= lng <= 180.0):
                logger.warning(f"Coordinates out of valid range: lat={lat}, lng={lng}")
                return None
        except (ValueError, TypeError) as e:
            logger.warning(f"Invalid coordinates: lat={lat}, lng={lng}, error={e}")
            return None

        # List of map services to try (in order of preference)
        zoom = 15
        map_services = [
            # OpenStreetMap static map (primary) - using staticmap.openstreetmap.de
            {
                "url": f"https://staticmap.openstreetmap.de/staticmap.php?center={lat},{lng}&zoom={zoom}&size=800x500&markers={lat},{lng},lightblue1",
                "name": "OpenStreetMap Static"
            },
            # Alternative: OpenStreetMap tile service (single tile)
            {
                "url": f"https://tile.openstreetmap.org/{zoom}/{int((lng + 180) / 360 * (2 ** zoom))}/{int((1 - math.log(math.tan(math.radians(lat)) + 1 / math.cos(math.radians(lat))) / math.pi) / 2 * (2 ** zoom))}.png",
                "name": "OpenStreetMap Tile"
            }
        ]

        # Try each service until one works
        for service in map_services:
            try:
                logger.info(f"Trying {service['name']} for map at {lat}, {lng}")

                # Use sync httpx client with proper headers and longer timeout
                with httpx.Client(timeout=15.0, follow_redirects=True) as client:
                    resp = client.get(
                        service["url"],
                        headers={
                            "User-Agent": "PropertyWalay-PPT-Generator/1.0 (https://propertywalay.com)",
                            "Accept": "image/png,image/jpeg,image/webp,*/*",
                            "Referer": "https://propertywalay.com"
                        }
                    )
                    resp.raise_for_status()

                    # Verify it's actually an image
                    content_type = resp.headers.get("content-type", "").lower()
                    if "image" in content_type and len(resp.content) > 100:
                        logger.info(f"Successfully fetched map from {service['name']} ({len(resp.content)} bytes)")
                        return resp.content
                    else:
                        logger.warning(f"Invalid image response from {service['name']}: content-type={content_type}, size={len(resp.content)}")
                        continue

            except httpx.TimeoutException:
                logger.warning(f"Map fetch from {service['name']} timed out")
                continue
            except httpx.ConnectError as e:
                # Log at debug level since fallback will be tried
                logger.debug(f"Connection error with {service['name']}: {e}")
                continue
            except httpx.HTTPStatusError as e:
                logger.warning(f"{service['name']} returned error {e.response.status_code}")
                continue
            except Exception as e:
                logger.warning(f"Failed to fetch map from {service['name']}: {type(e).__name__}: {e}")
                continue

        # All services failed
        logger.warning(f"All map services failed for coordinates {lat}, {lng}")
        return None


class ImageOptimizer:
    """Handles image downloading and optimization for PPT"""

    MAX_WIDTH = 1920
    QUALITY = 85
    TIMEOUT = 10

    @staticmethod
    async def fetch_and_optimize(url: str) -> Optional[bytes]:
        """
        Download and optimize an image from URL.

        Args:
            url: Image URL to download

        Returns:
            Optimized image bytes or None if failed
        """
        try:
            async with httpx.AsyncClient() as client:
                response = await client.get(url, timeout=ImageOptimizer.TIMEOUT)
                response.raise_for_status()

            # Open and optimize
            img = Image.open(io.BytesIO(response.content))

            # Resize if too large
            if img.width > ImageOptimizer.MAX_WIDTH:
                ratio = ImageOptimizer.MAX_WIDTH / img.width
                new_size = (ImageOptimizer.MAX_WIDTH, int(img.height * ratio))
                img = img.resize(new_size, Image.Resampling.LANCZOS)

            # Convert to RGB if needed
            if img.mode in ('RGBA', 'LA', 'P'):
                # Create white background for transparency
                background = Image.new('RGB', img.size, (255, 255, 255))
                if img.mode == 'P':
                    img = img.convert('RGBA')
                background.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else None)
                img = background
            elif img.mode != 'RGB':
                img = img.convert('RGB')

            # Save optimized
            output = io.BytesIO()
            img.save(output, format='JPEG', quality=ImageOptimizer.QUALITY, optimize=True)
            output.seek(0)

            return output.getvalue()

        except Exception as e:
            logger.warning(f"Failed to fetch/optimize image {url}: {e}")
            return None

    @staticmethod
    async def fetch_multiple(urls: List[str], max_images: int = 5) -> List[bytes]:
        """
        Fetch and optimize multiple images concurrently.

        Args:
            urls: List of image URLs
            max_images: Maximum number of images to fetch

        Returns:
            List of optimized image bytes
        """
        import asyncio

        urls = urls[:max_images]
        tasks = [ImageOptimizer.fetch_and_optimize(url) for url in urls]
        results = await asyncio.gather(*tasks, return_exceptions=True)

        # Filter out None and exceptions
        return [r for r in results if isinstance(r, bytes)]
