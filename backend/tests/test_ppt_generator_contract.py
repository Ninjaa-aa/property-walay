import io
import sys
import unittest
from pathlib import Path
from typing import Any, Dict, Optional

from PIL import Image
from pptx import Presentation


BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.services.ppt.ppt_generator import PropertyPPTGenerator  # noqa: E402


def _make_test_image_bytes(color: tuple[int, int, int] = (30, 90, 160)) -> bytes:
    """Create a small JPEG in-memory for pptx insertion."""
    img = Image.new("RGB", (240, 140), color)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def _base_property_data() -> Dict[str, Any]:
    return {
        "our_id": "123e4567-e89b-12d3-a456-426614174000",
        "title": "Test Property",
        "prop_type": "Apartment",
        "prop_subtype": "Residential",
        "area_size": 120.0,
        "area_unit": "sqft",
        "beds": 2,
        "baths": 2,
        "area_name": "Lahore",
        "current_price": 5000000.0,
        "currency": "PKR",
        "source": "graana",
        "source_human_id": "GRA-00000123",
        "poc_name": "Agent Name",
        "poc_number": "+92 300 0000000",
        "scraped_features": {},
        "scraped_amenities": {},
    }


class TestPptGeneratorContract(unittest.TestCase):
    def test_base_slides_no_optional_data(self):
        generator = PropertyPPTGenerator(template_name="Clean Minimal")

        property_data = _base_property_data()
        # Explicitly keep optional sections off
        property_data["scraped_features"] = {}
        property_data["scraped_amenities"] = {}
        property_data["latitude"] = None
        property_data["longitude"] = None

        ppt_stream = generator.generate(property_data, images=None)
        ppt_stream.seek(0)
        prs = Presentation(ppt_stream)

        # Title + Overview + Details + Contact
        self.assertEqual(len(prs.slides), 4)

    def test_features_slide_added_for_2_sections(self):
        generator = PropertyPPTGenerator(template_name="Clean Minimal")

        property_data = _base_property_data()
        property_data["scraped_features"] = {
            "Primary": ["Balcony", "Parking"],
            "Utilities": ["Gas", "Electricity"],
        }
        property_data["scraped_amenities"] = {}
        property_data["latitude"] = None
        property_data["longitude"] = None

        ppt_stream = generator.generate(property_data, images=None)
        ppt_stream.seek(0)
        prs = Presentation(ppt_stream)

        # Base 4 + 1 features slide (2 sections => 1 chunk => 1 slide)
        self.assertEqual(len(prs.slides), 5)

    def test_gallery_slide_added_when_images_len_gt_1(self):
        generator = PropertyPPTGenerator(template_name="Clean Minimal")

        property_data = _base_property_data()
        property_data["scraped_features"] = {}
        property_data["scraped_amenities"] = {}
        property_data["latitude"] = None
        property_data["longitude"] = None

        img1 = _make_test_image_bytes((40, 120, 80))
        img2 = _make_test_image_bytes((120, 60, 40))

        ppt_stream = generator.generate(property_data, images=[img1, img2])
        ppt_stream.seek(0)
        prs = Presentation(ppt_stream)

        # Base 4 + 1 gallery slide
        self.assertEqual(len(prs.slides), 5)

    def test_location_slide_added_when_lat_lng_present(self):
        generator = PropertyPPTGenerator(template_name="Clean Minimal")

        property_data = _base_property_data()
        property_data["scraped_features"] = {}
        property_data["scraped_amenities"] = {}
        property_data["latitude"] = 31.0
        property_data["longitude"] = 74.0

        # Avoid network: static map fetch returns None => placeholder still renders.
        generator._fetch_static_map = lambda *_args, **_kwargs: None  # type: ignore[attr-defined]

        ppt_stream = generator.generate(property_data, images=None)
        ppt_stream.seek(0)
        prs = Presentation(ppt_stream)

        # Base 4 + 1 location slide
        self.assertEqual(len(prs.slides), 5)


if __name__ == "__main__":
    unittest.main()

